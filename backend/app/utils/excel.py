import re
from typing import Optional
from io import BytesIO

from openpyxl import load_workbook


def parse_excel_file(file_content: bytes) -> dict:
    """Parse an Excel file and return structured data with validation info.
    
    Returns:
        {
            "headers": [...],
            "rows": [...],
            "total_rows": int,
            "error": str or None
        }
    """
    try:
        wb = load_workbook(filename=BytesIO(file_content), read_only=True, data_only=True)
        ws = wb.active

        headers = []
        rows = []

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(cell).strip() if cell else "" for cell in row]
                continue
            row_data = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    row_data[headers[j]] = str(cell).strip() if cell is not None else ""
            # Skip completely empty rows
            if any(v for v in row_data.values()):
                rows.append(row_data)

        wb.close()

        return {
            "headers": headers,
            "rows": rows,
            "total_rows": len(rows),
            "error": None,
        }
    except Exception as e:
        return {
            "headers": [],
            "rows": [],
            "total_rows": 0,
            "error": str(e),
        }


def validate_usn(usn: str) -> bool:
    """Validate USN format - basic alphanumeric check."""
    if not usn or len(usn.strip()) < 3:
        return False
    # Allow alphanumeric USNs (letters, digits, hyphens, underscores)
    return bool(re.match(r'^[A-Za-z0-9_\-]+$', usn.strip()))


def detect_column_mapping(headers: list) -> dict:
    """Auto-detect column mappings from header names.
    
    Returns: {"usn": col_name, "name": col_name, "class": col_name}
    """
    mapping = {"usn": None, "name": None, "class": None}

    usn_patterns = ["usn", "usn no", "usn number", "roll no", "roll number", "registration", "reg no", "student id", "id"]
    name_patterns = ["name", "student name", "full name", "student"]
    class_patterns = ["class", "section", "branch", "department", "dept", "class name", "batch"]

    for header in headers:
        h_lower = header.lower().strip()
        if not mapping["usn"] and h_lower in usn_patterns:
            mapping["usn"] = header
        elif not mapping["name"] and h_lower in name_patterns:
            mapping["name"] = header
        elif not mapping["class"] and h_lower in class_patterns:
            mapping["class"] = header

    return mapping


def validate_import_data(
    rows: list,
    usn_column: str,
    name_column: Optional[str] = None,
    class_column: Optional[str] = None,
    existing_usns: set = None,
) -> dict:
    """Validate import data and return validation results.
    
    Returns:
        {
            "valid": [...],
            "duplicates": [...],
            "invalid": [...],
            "existing": [...],
            "summary": {...}
        }
    """
    if existing_usns is None:
        existing_usns = set()

    valid = []
    duplicates = []
    invalid = []
    existing = []
    seen_usns = set()

    for i, row in enumerate(rows):
        usn = row.get(usn_column, "").strip().upper()
        name = row.get(name_column, "").strip() if name_column and name_column in row else ""
        class_name = row.get(class_column, "").strip() if class_column and class_column in row else ""

        row_num = i + 2  # Excel row number (1-indexed header + data)

        if not usn:
            invalid.append({
                "row": row_num,
                "usn": usn,
                "name": name,
                "class": class_name,
                "reason": "Empty or missing USN"
            })
            continue

        if not validate_usn(usn):
            invalid.append({
                "row": row_num,
                "usn": usn,
                "name": name,
                "class": class_name,
                "reason": "Invalid USN format"
            })
            continue

        if usn in seen_usns:
            duplicates.append({
                "row": row_num,
                "usn": usn,
                "name": name,
                "class": class_name,
                "reason": "Duplicate USN in file"
            })
            continue

        if usn in existing_usns:
            existing.append({
                "row": row_num,
                "usn": usn,
                "name": name,
                "class": class_name,
                "reason": "Already registered in system"
            })
            continue

        seen_usns.add(usn)
        valid.append({
            "row": row_num,
            "usn": usn,
            "name": name,
            "class": class_name,
        })

    return {
        "valid": valid,
        "duplicates": duplicates,
        "invalid": invalid,
        "existing": existing,
        "summary": {
            "total_rows": len(rows),
            "valid_count": len(valid),
            "duplicate_count": len(duplicates),
            "invalid_count": len(invalid),
            "existing_count": len(existing),
        }
    }
